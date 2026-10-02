#!/usr/bin/env python3
"""Kargo PM / SPM CV screening pipeline.

  python screen.py                      # screen every CV in ./cvs for PM and SPM
  python screen.py --roles PM           # one role only
  python screen.py --rescore            # re-score every CV, not just new or changed ones
  python screen.py --calibrate          # compare results.csv with ratings.csv
  python report.py                      # rebuild report.html from the last run

The model returns per-criterion scores with exact quotes from the CV. Every
quote is checked against the CV text; a score above 0 with no surviving quote
is reset to 0. Weighted score, gate, rescue rules, flags and route are all
computed in rubric.py, never by the model.
"""
import argparse
import csv
import json
import os
import re
import sys
from pathlib import Path
from typing import Dict, List, Literal, Optional, Tuple

from pydantic import BaseModel

import rubric
from rubric import ROLES, Facts, Role

DEFAULT_MODEL = "claude-sonnet-5"          # used when ANTHROPIC_API_KEY is the key set
DEFAULT_GEMINI_MODEL = "gemini-3.6-flash"   # used when GEMINI_API_KEY is set (takes priority)
GEMINI_RETRIES = 3
MIN_TEXT_CHARS = 400  # below this the parse is treated as low confidence (R4)
MIN_QUOTE_CHARS = 8

# Evidence that is only a credential is never scoreable (rubric P3 / never-score list).
CREDENTIAL_RE = re.compile(
    r"\b(mba|pgdm|bba|b\.\s?(com|e|tech|sc)|m\.\s?sc|certif\w*|"
    r"reforge|pmp|cspo|iim\w*|xlri|iit\w*|universit\w*|college|alumn\w*)\b",
    re.I,
)
OUTCOME_RE = re.compile(r"\d|adopt|launch|shipp|kill|reduc|increas|cut|migrat|resolv|promot", re.I)


# ---------------------------------------------------------------- extraction

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\w)\+?\d[\d\s().-]{8,}\d(?!\w)")


def redact(text: str) -> str:
    """Strip contact details before the CV is sent to the model (they never score)."""
    return PHONE_RE.sub("[phone]", EMAIL_RE.sub("[email]", text))


def extract_text(path: Path) -> str:
    return extract_text_bytes(path.name, path.read_bytes())


def extract_text_bytes(filename: str, data: bytes) -> str:
    import io

    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    raise ValueError(f"unsupported file type: {filename}")


def normalize(s: str) -> str:
    s = s.lower()
    s = s.translate(str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"',
                                   "–": "-", "—": "-", "−": "-", " ": " "}))
    s = re.sub(r"[•●▪‣]", " ", s)  # bullets
    return re.sub(r"\s+", " ", s).strip()


def quote_in_cv(quote: str, norm_cv: str) -> bool:
    q = normalize(quote).strip(" .\"'")
    return len(q) >= MIN_QUOTE_CHARS and q in norm_cv


def is_credential_only(quote: str) -> bool:
    return bool(CREDENTIAL_RE.search(quote)) and not OUTCOME_RE.search(quote)


# ------------------------------------------------------------ model contract

class CriterionResult(BaseModel):
    id: str
    score: int
    evidence: List[str]
    reason: str


class Assessment(BaseModel):
    candidate_name: str
    criteria: List[CriterionResult]
    g1_pass: bool
    g1_quote: str
    pm_title_years: float
    ops_tenure_years: float
    ownership_role_years: float
    mumbai_in_office_stated: bool
    mumbai_quote: str
    parse_confidence: Literal["high", "low"]
    brief: str


def system_prompt(role: Role) -> str:
    crit = "\n".join(
        f"  {c.id}  {c.name}  (weight {c.weight:g}%)\n      {c.anchor}" for c in role.criteria
    )
    return f"""You screen CVs for Kargo, a Series A freight/3PL SaaS company in Mumbai, \
for the {role.title} role. You score evidence, not titles.

{rubric.SCORE_ANCHORS}

Criteria for {role.title}:
{crit}

Rules for evidence:
- For each criterion return 0-4 evidence strings copied EXACTLY, character for character, \
from the CV text (a phrase or sentence, not a paraphrase). Code checks every quote against \
the CV and discards any that do not match; a score above 0 with no matching quote becomes 0.
- No quote means score 0. Adjectives and self-descriptions ("results-driven", "passionate") score 0.
- NEVER give points for: {", ".join(rubric.NEVER_SCORE)}. Score the work, not where it was \
done or what qualifies the person.
- reason: one line explaining the score against the anchor.
- Return all six criteria ({", ".join(role.ids)}), each exactly once.

Also extract these facts (they drive gates and flags, not the score):
- g1_pass / g1_quote: {rubric.G1_DEFINITION} g1_quote is an exact CV quote proving it, \
or "" if g1_pass is false.
- pm_title_years: total years holding a Product Manager title (any level). 0 if none.
- ops_tenure_years: total years in a paid, ground-level logistics / freight / 3PL / \
customs / port / supply-chain operations job (doing the work, not selling or building for it).
- ownership_role_years: total years in roles where the person owned a product, module, \
platform, function or P&L decision end-to-end.
- mumbai_in_office_stated / mumbai_quote: true only if the CV says the person is based in \
Mumbai or will relocate there; mumbai_quote is the exact CV text, or "".
- parse_confidence: "low" if the text looks garbled, truncated, or not a standard CV; \
otherwise "high".
- candidate_name: the candidate's name as written on the CV.
- brief: 2-3 sentences for the hiring manager, from the CV only: the strongest evidence for \
this role, the main gap against it, and one specific thing to probe in interview. It is \
context for a human and does not affect the score."""


class ModelError(RuntimeError):
    """Any model/provider failure. The CV is routed to HUMAN REVIEW (CLI) or not saved (web)."""


def make_client(model: Optional[str] = None) -> Tuple[object, str]:
    """Pick the provider from whichever key is set. GEMINI_API_KEY wins if both are."""
    gemini_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if gemini_key:
        from google import genai

        return genai.Client(api_key=gemini_key), model or os.environ.get("GEMINI_MODEL", DEFAULT_GEMINI_MODEL)
    if os.environ.get("ANTHROPIC_API_KEY", "").strip():
        import anthropic

        return anthropic.Anthropic(), model or os.environ.get("KARGO_MODEL", DEFAULT_MODEL)
    raise ModelError("No model key is set. Add GEMINI_API_KEY (or ANTHROPIC_API_KEY).")


def assess(client, model: str, role: Role, cv_text: str) -> Assessment:
    if model.startswith("gemini"):
        return _assess_gemini(client, model, role, cv_text)
    return _assess_claude(client, model, role, cv_text)


def _assess_gemini(client, model: str, role: Role, cv_text: str) -> Assessment:
    import time

    from google.genai import errors, types

    config = types.GenerateContentConfig(
        system_instruction=system_prompt(role),
        response_mime_type="application/json",
        response_schema=Assessment,
    )
    for attempt in range(GEMINI_RETRIES + 1):
        try:
            response = client.models.generate_content(
                model=model, contents=f"<cv>\n{cv_text}\n</cv>", config=config)
            break
        except errors.APIError as e:
            if "PerDay" in str(e):
                raise ModelError("Gemini's daily free-tier quota is used up (resets midnight Pacific). "
                                 "A key with billing enabled removes the limit.") from e
            if e.code not in (429, 500, 503) or attempt == GEMINI_RETRIES:
                raise ModelError(f"Gemini error {e.code}: {e.message or e}") from e
            hinted = re.search(r"retry in ([\d.]+)s", str(e))
            time.sleep(min(20.0, float(hinted.group(1)) + 1 if hinted else 4 * 2 ** attempt))
    parsed = response.parsed
    if isinstance(parsed, Assessment):
        return parsed
    try:
        return Assessment.model_validate_json(response.text or "")
    except ValueError as e:
        reason = getattr((response.candidates or [None])[0], "finish_reason", None)
        raise ModelError(f"Gemini returned no usable JSON (finish reason: {reason})") from e


def _assess_claude(client, model: str, role: Role, cv_text: str) -> Assessment:
    import anthropic

    try:
        return _claude_call(client, model, role, cv_text)
    except anthropic.APIError as e:
        raise ModelError(f"Claude error: {e}") from e


def _claude_call(client, model: str, role: Role, cv_text: str) -> Assessment:
    response = client.messages.parse(
        model=model,
        max_tokens=16000,
        system=[{"type": "text", "text": system_prompt(role),
                 "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": f"<cv>\n{cv_text}\n</cv>"}],
        output_format=Assessment,
        output_config={"effort": "high"},
    )
    if response.stop_reason == "refusal":
        raise RuntimeError(f"model refused: {response.stop_details}")
    if response.stop_reason == "max_tokens":
        raise RuntimeError("model output truncated at max_tokens")
    if response.parsed_output is None:
        raise RuntimeError("model returned no parseable output")
    return response.parsed_output


# ------------------------------------------------------------- verification

def verify(role: Role, a: Assessment, cv_text: str) -> Tuple[Dict[str, int], Dict, Facts]:
    """Apply the quote checks and turn the model output into scores + facts."""
    norm_cv = normalize(cv_text)
    by_id = {c.id: c for c in a.criteria}
    scores: Dict[str, int] = {}
    evidence: Dict[str, dict] = {}

    for cid in role.ids:
        c = by_id.get(cid)
        if c is None:
            scores[cid] = 0
            evidence[cid] = {"score": 0, "model_score": None, "evidence": [], "dropped": [],
                             "reason": "criterion missing from model output"}
            continue
        kept, dropped = [], []
        for q in c.evidence:
            if not quote_in_cv(q, norm_cv):
                dropped.append({"quote": q, "why": "not found in CV"})
            elif is_credential_only(q):
                dropped.append({"quote": q, "why": "credential (never-score list)"})
            else:
                kept.append(q)
        score = max(0, min(4, c.score))
        note = c.reason
        if score > 0 and not kept:
            score, note = 0, f"reset to 0: no verified quote (model said {c.score}: {c.reason})"
        scores[cid] = score
        evidence[cid] = {"score": score, "model_score": c.score, "evidence": kept,
                         "dropped": dropped, "reason": note}

    g1 = a.g1_pass and quote_in_cv(a.g1_quote, norm_cv)
    mumbai = a.mumbai_in_office_stated and quote_in_cv(a.mumbai_quote, norm_cv)
    low_text = len(cv_text.strip()) < MIN_TEXT_CHARS
    facts = Facts(
        g1_pass=g1,
        pm_title_years=max(0.0, a.pm_title_years),
        ops_tenure_years=max(0.0, a.ops_tenure_years),
        ownership_role_years=max(0.0, a.ownership_role_years),
        mumbai_in_office_stated=mumbai,
        parse_confidence="low" if (low_text or a.parse_confidence == "low") else "high",
    )
    evidence["_facts"] = {
        "g1_pass": g1, "g1_quote": a.g1_quote if g1 else "",
        "g1_model_said": a.g1_pass,
        "pm_title_years": facts.pm_title_years,
        "ops_tenure_years": facts.ops_tenure_years,
        "ownership_role_years": facts.ownership_role_years,
        "mumbai_in_office_stated": mumbai, "mumbai_quote": a.mumbai_quote if mumbai else "",
        "parse_confidence": facts.parse_confidence, "text_chars": len(cv_text),
    }
    evidence["_brief"] = a.brief.strip()
    return scores, evidence, facts


# --------------------------------------------------------------------- run

ALL_IDS = [cid for r in ROLES.values() for cid in r.ids]
RESULT_FIELDS = (["candidate", "file", "role"] + ALL_IDS +
                 ["total", "g1", "route", "route_reason", "rules_fired", "flags", "error"])


def load_previous(out_dir: Path) -> Tuple[Dict[Tuple[str, str], dict], Dict[str, dict]]:
    """Rows and evidence from the last run, keyed so unchanged CVs can be reused."""
    rows: Dict[Tuple[str, str], dict] = {}
    evidence: Dict[str, dict] = {}
    try:
        with open(out_dir / "results.csv") as f:
            rows = {(r["file"], r["role"]): r for r in csv.DictReader(f)}
        evidence = json.loads((out_dir / "evidence.json").read_text())
    except (OSError, ValueError, KeyError):
        return {}, {}
    return rows, evidence


def text_sha(text: str) -> str:
    import hashlib

    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def score_cv(client, model: str, filename: str, text: str, roles: List[Role],
             extract_error: str = "", log=print) -> Tuple[List[dict], dict]:
    """Score one CV's text for each role. Returns (result rows, evidence for this file).

    Shared by the CLI and the web API, so both score identically. Roles are
    scored in parallel; any failure routes that role to HUMAN REVIEW.
    """
    from concurrent.futures import ThreadPoolExecutor

    text = redact(text)
    stem = Path(filename).stem
    file_ev: dict = {"candidate": stem, "_sha": text_sha(text)}

    def one(role: Role) -> Tuple[dict, Optional[dict]]:
        row = {"file": filename, "role": role.key, "candidate": stem, "error": extract_error}
        if extract_error or not text.strip():
            row.update(route=rubric.REVIEW, route_reason="R4 unreadable CV",
                       rules_fired="R4", error=extract_error or "no text extracted")
            return row, None
        try:
            a = assess(client, model, role, text)
        except RuntimeError as e:  # ModelError, refusals, truncation
            row.update(route=rubric.REVIEW, route_reason="model error", error=str(e))
            return row, None
        scores, ev, facts = verify(role, a, text)
        d = rubric.decide(role, scores, facts)
        row.update(candidate=a.candidate_name.strip() or stem, **scores,
                   total=f"{d.total:.2f}", g1="Y" if d.g1_pass else "N",
                   route=d.route, route_reason=d.reason,
                   rules_fired=";".join(d.rules_fired), flags=";".join(d.flags))
        return row, ev

    with ThreadPoolExecutor(max_workers=len(roles) or 1) as pool:
        results = list(pool.map(one, roles))

    rows = []
    for role, (row, ev) in zip(roles, results):
        rows.append(row)
        if ev is not None:
            file_ev[role.key] = ev
            file_ev["candidate"] = row["candidate"]
        log(f"  {filename} [{role.key}] ... "
            + (f"{float(row['total']):6.2f}  {row['route']}" if row.get("total") else
               f"{row['route']}  ({row.get('error') or row.get('route_reason')})"))
    return rows, file_ev


def screen(cv_dir: Path, out_dir: Path, roles: List[Role], model: str, rescore: bool = False,
           client=None) -> int:
    files = sorted(p for p in cv_dir.iterdir()
                   if p.suffix.lower() in (".docx", ".pdf") and not p.name.startswith("~$"))
    if not files:
        print(f"No .docx/.pdf files in {cv_dir}", file=sys.stderr)
        return 1

    if client is None:
        client, model = make_client(model)
    rows, evidence_out = [], {}
    prev_rows, prev_ev = ({}, {}) if rescore else load_previous(out_dir)
    new_calls = reused = 0

    for path in files:
        try:
            text, err = extract_text(path), ""
        except Exception as e:  # unreadable file: route to a human, don't drop it
            text, err = "", f"extract failed: {e}"
        sha = text_sha(redact(text))
        old_ev = prev_ev.get(path.name, {})
        todo = []
        for role in roles:
            old_row = prev_rows.get((path.name, role.key))
            if (old_row and not old_row.get("error") and old_ev.get("_sha") == sha
                    and role.key in old_ev):
                rows.append(old_row)
                entry = evidence_out.setdefault(path.name, {"candidate": old_row["candidate"], "_sha": sha})
                entry[role.key] = old_ev[role.key]
                reused += 1
                print(f"  {path.name} [{role.key}] ... already scored, kept ({old_row['route']})")
            else:
                todo.append(role)
        if not todo:
            continue
        new_rows, file_ev = score_cv(client, model, path.name, text, todo, err)
        new_calls += len(todo)
        rows.extend(new_rows)
        entry = evidence_out.setdefault(path.name, {"_sha": sha})
        entry.update({k: v for k, v in file_ev.items()})

    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "results.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=RESULT_FIELDS)
        w.writeheader()
        w.writerows(rows)
    with open(out_dir / "evidence.json", "w") as f:
        json.dump(evidence_out, f, indent=2, ensure_ascii=False)
    import report

    report.render(rows, evidence_out, out_dir / "report.html", model)
    print(f"\n{new_calls} scored now, {reused} kept from earlier runs"
          + (" (use --rescore to redo them)" if reused else ""))
    print(f"Wrote {out_dir / 'results.csv'}, {out_dir / 'evidence.json'} and {out_dir / 'report.html'}")
    return 0


# --------------------------------------------------------------- calibrate

def _ranks(xs: List[float]) -> List[float]:
    order = sorted(range(len(xs)), key=lambda i: xs[i])
    ranks = [0.0] * len(xs)
    i = 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and xs[order[j + 1]] == xs[order[i]]:
            j += 1
        for k in range(i, j + 1):
            ranks[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return ranks


def spearman(xs: List[float], ys: List[float]) -> Optional[float]:
    if len(xs) < 3:
        return None
    rx, ry = _ranks(xs), _ranks(ys)
    mx, my = sum(rx) / len(rx), sum(ry) / len(ry)
    cov = sum((a - mx) * (b - my) for a, b in zip(rx, ry))
    vx = sum((a - mx) ** 2 for a in rx) ** 0.5
    vy = sum((b - my) ** 2 for b in ry) ** 0.5
    return None if vx == 0 or vy == 0 else cov / (vx * vy)


def name_key(s: str) -> str:
    return re.sub(r"[^a-z]", "", s.lower())


def calibrate(results_path: Path, ratings_path: Path) -> int:
    if not results_path.exists():
        print(f"{results_path} not found - run screen.py first", file=sys.stderr)
        return 1
    if not ratings_path.exists():
        print(f"{ratings_path} not found (columns: name,role,actual_rating)", file=sys.stderr)
        return 1
    with open(results_path) as f:
        results = {(name_key(r["candidate"]), r["role"].upper()): r
                   for r in csv.DictReader(f) if r.get("total")}
    with open(ratings_path) as f:
        ratings = list(csv.DictReader(f))

    for role in ROLES:
        pairs, misses, unmatched = [], [], []
        for r in ratings:
            if r["role"].strip().upper() != role:
                continue
            res = results.get((name_key(r["name"]), role))
            if res is None:
                unmatched.append(r["name"])
                continue
            rating = float(r["actual_rating"])
            pairs.append((float(res["total"]), rating, r["name"], res["route"]))
            if rating >= 4 and res["route"] == rubric.REJECT:
                misses.append(f"{r['name']} (rating {rating:g}, score {res['total']})")
        if not pairs and not unmatched:
            continue
        rho = spearman([p[0] for p in pairs], [p[1] for p in pairs])
        print(f"\n== {role}  (n={len(pairs)})")
        print(f"Spearman rho, rubric score vs rating: "
              f"{'n/a (need 3+ rows with variance)' if rho is None else f'{rho:+.3f}'}")
        for score, rating, name, route in sorted(pairs, key=lambda p: -p[0]):
            print(f"  {score:6.2f}  rating {rating:g}  {route:<13} {name}")
        print("Rated 4+ but REJECTED: " + (", ".join(misses) if misses else "none"))
        if unmatched:
            print("No screening result for: " + ", ".join(unmatched))
    return 0


def main(argv=None) -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cvs", type=Path, default=here / "cvs")
    ap.add_argument("--out", type=Path, default=here)
    ap.add_argument("--roles", nargs="+", choices=list(ROLES), default=list(ROLES))
    ap.add_argument("--model", default=None,
                    help=f"default: {DEFAULT_GEMINI_MODEL} with GEMINI_API_KEY, else {DEFAULT_MODEL}")
    ap.add_argument("--calibrate", action="store_true",
                    help="compare results.csv with ratings.csv instead of screening")
    ap.add_argument("--ratings", type=Path, default=here / "ratings.csv")
    ap.add_argument("--rescore", action="store_true",
                    help="re-score every CV, not just new or changed ones")
    args = ap.parse_args(argv)

    if args.calibrate:
        return calibrate(args.out / "results.csv", args.ratings)
    try:
        client, model = make_client(args.model)
    except ModelError as e:
        print(e, file=sys.stderr)
        return 1
    print(f"Model: {model}")
    return screen(args.cvs, args.out, [ROLES[k] for k in args.roles], model, args.rescore, client)


if __name__ == "__main__":
    sys.exit(main())
